"""Numeric and structural audit of MANUSCRIPT_FRAMEWORK_ML_V11B_TEMPLATE_E.docx (read-only; writes one NEW markdown file).

1. Re-opens the docx with python-docx and extracts every paragraph and table.
2. Confirms each registered block text (NUMERIC_REGISTRY_FRAMEWORK_V11B.json) is present in the docx.
3. Extracts every number token from the docx text and matches it to a registered source in the same block.
4. Re-reads every CSV-backed value from its file and row and compares at printed precision.
5. Structural and content checks, including that no value of the earlier surrogate run appears.

usage: python audit_numbers_framework_v11b.py <package_root> <python-docx lib dir> [folder holding docx+registry]
"""
import json
import os
import re
import sys
from pathlib import Path

import pandas as pd

ROOT = Path(sys.argv[1]).resolve()
sys.dont_write_bytecode = True
sys.path.insert(0, sys.argv[2])
from docx import Document  # noqa: E402

DIR = Path(sys.argv[3]).resolve() if len(sys.argv) > 3 else ROOT / "manuscript"
DOCX = DIR / "MANUSCRIPT_FRAMEWORK_ML_V11B_TEMPLATE_E.docx"
REG = DIR / "NUMERIC_REGISTRY_FRAMEWORK_V11B.json"
OUT = DIR / "NUMERIC_AUDIT_FRAMEWORK_V11B.md"
REPLACE = os.environ.get("NEGI_V11_REPLACE_OWN_OUTPUT") == "1"
if OUT.exists() and not REPLACE:
    raise SystemExit(f"Refusing to overwrite {OUT}")

MINUS = "−"
TOKEN = re.compile(r"(?<![\w.])\d[\d,]*(?:\.\d+)?")
STRIP = [
    r"\[\d+(?:[,–]\d+)*\]",
    r"Eqs?\.\s*\(\d+\)(?:–\(\d+\))?",
    r"Tables?\sS?\d+", r"Figures?\s\d+[abc]?", r"Sections?\s\d+(?:\.\d+)*",
    r"RQ\d(?:\sand\sRQ\d)?", r"Property\s\d", r"[Ss]tages?\s\d(?:\sto\s\d)?", r"[Gg]ates?\s\d(?:\sand\s\d)?",
    r"Model\s[AB]",
]
SKIP_ROLES = ("reference", "supplementary list", "references heading", "appendix heading", "front matter", "keywords")


def tokens(text):
    for pat in STRIP:
        text = re.sub(pat, " ", text)
    return [t.replace(",", "").rstrip(".") for t in TOKEN.findall(text)]


def fmt(x, d):
    s = f"{abs(x):,.{d}f}"
    return (MINUS if x < 0 and float(s.replace(",", "")) != 0 else "") + s


doc = Document(DOCX)
reg = json.loads(REG.read_text(encoding="utf-8"))
pars = [(p.style.name, p.text) for p in doc.paragraphs]
doc_pars = [t for _, t in pars]
doc_tables = [" | ".join(c.text for c in t.rows[0].cells) + " || " +
              " || ".join(" | ".join(c.text for c in r.cells) for r in t.rows[1:]) for t in doc.tables]
doc_text_set = set(doc_pars) | set(doc_tables)

lines, flags, csv_checks, literals = [], [], 0, []
n_tokens = n_matched = 0
_csv = {}


def csv(rel):
    if rel not in _csv:
        _csv[rel] = pd.read_csv(ROOT / rel)
    return _csv[rel]


block_text = {}
for b in reg["blocks"]:
    block_text[b["index"]] = block_text.get(b["index"], "") + "\n" + b["text"]
lines.append("| Block | Role | Number in docx | Source | Check |")
lines.append("|---|---|---|---|---|")
for b in reg["blocks"]:
    role, text = b["role"], b["text"]
    if text not in doc_text_set:
        flags.append(f"block {b['index']} ({role}) text not found in docx: {text[:70]}")
        continue
    if role in SKIP_ROLES:
        continue
    body = text
    if role == "heading":
        body = re.sub(r"^\d+(?:\.\d+)*\.\s", "", body)
    if role in ("table caption", "figure caption"):
        body = re.sub(r"^(Table|Figure)\s\d+\.\s", "", body)
    if role == "equation":
        body = re.sub(r"\(\d+\)\s*$", "", body)
    if role == "table cells":
        body = re.sub(r"(^|\|\|?\s)\d\.\s", r"\1", body)
    pool = {}
    for e in b["numbers"]:
        if e["kind"] == "csv":
            now = float(csv(e["file"]).iloc[e["row"] - 2][e["col"]])
            again = fmt(now * e.get("scale", 1.0), e["decimals"])
            csv_checks += 1
            ok = again == e["text"].lstrip("+")
            src = f"{e['file']} : row {e['row']} : {e['col']} = {now!r}" + (f" (x{e['scale']:g})" if e.get("scale", 1.0) != 1.0 else "")
            if not ok:
                flags.append(f"MISMATCH block {b['index']}: printed {e['text']} but {src} formats to {again}")
        elif e["kind"] == "csv_sum":
            now = float(sum(csv(e["file"]).iloc[r - 2][e["col"]] for r in e["rows"]))
            csv_checks += 1
            ok = fmt(now, 0) == e["text"]
            src = f"{e['file']} : rows {e['rows']} : sum of {e['col']} = {now!r}"
            if not ok:
                flags.append(f"MISMATCH block {b['index']}: printed {e['text']} but {src}")
        else:
            ok, src = True, e["source"]
        if e["text"] not in block_text[b["index"]]:
            flags.append(f"block {b['index']}: registered value {e['text']} is not in the rendered text")
        for tk in tokens(e["text"]):
            pool.setdefault(tk, []).append((src, ok, e["kind"]))
    for tk in tokens(body):
        n_tokens += 1
        if tk in pool:
            n_matched += 1
            src, ok, kind = pool[tk][0]
            status = {"csv": "matches table at printed precision", "csv_sum": "matches table sum",
                      "const": "constant; source stated"}[kind]
            lines.append(f"| {b['index']} | {role} | {tk} | {src} | {status if ok else 'MISMATCH'} |")
        elif re.fullmatch(r"\d", tk):
            literals.append((b["index"], role, tk))
        else:
            flags.append(f"UNSOURCED number {tk} in block {b['index']} ({role}): {text[:80]}")

# ---------------- structural and content checks ----------------
full = "\n".join(doc_pars + doc_tables)
struct = []


def check(name, ok, detail=""):
    struct.append((name, "pass" if ok else "FAIL", detail))
    if not ok:
        flags.append(f"STRUCTURE: {name} {detail}")


check("no backticks", "`" not in full)
check('no "Net Energy Gain Index"', "Net Energy Gain Index" not in full)
check('no fractional-vegetation-cover thresholds ("5% FVC", "15-20%", "energy-optimal")',
      not re.search(r"FVC|15[–-]20\s?%|energy-optimal|optimal vegetation", full))
# "0.005" alone is not listed: the v11 calendar check has a paired change of -0.005 (gee_common_month_paired_v6.csv).
OLD = ["0.543", "0.547", "0.538", "1.74", "27,234", "27234", "0.91 ", "0.005 at 99", "0.304", "0.038", "69.5", "0.817", "3.64"]
non_ref = "\n".join(t for s, t in pars if s != "MDPI_7.1_References") + "\n" + "\n".join(doc_tables)
hits = [o for o in OLD if o in non_ref]
check("no value of the earlier surrogate run appears outside the reference list", not hits, ", ".join(hits))
check("no template instruction text left",
      not re.search(r"keyword 1|Subsection|This is a figure|entry 1|should briefly place|First bullet", full))
used = sorted({s for s, t in pars if t.strip()})
check("only template styles are used", all(s.startswith("MDPI_") or s in ("Normal", "p1") for s in used), ", ".join(used))


def hlevel(style):
    m = re.search(r"heading(\d)$", style)
    return int(m.group(1)) if m else None


items = [(s, t.strip()) for s, t in pars]
empty = []
for i, (sty, txt) in enumerate(items):
    lvl = hlevel(sty)
    if lvl and txt != "References":
        nxt = next(((s, t) for s, t in items[i + 1:] if t), None)
        if not txt or nxt is None or (hlevel(nxt[0]) and hlevel(nxt[0]) <= lvl):
            empty.append(txt or "<blank>")
check("no empty headings", not empty, "; ".join(empty))
heads = [t for s, t in items if hlevel(s) and t and t != "References"]
check("headings numbered in the template form (1. / 1.1.)", all(re.match(r"^\d+(\.\d+)?\.\s\S", h) for h in heads), f"{len(heads)} headings")
top = [h.split(". ", 1)[1] for h in heads if re.match(r"^\d+\.\s", h)]
check("top-level sections follow the template", top == ["Introduction", "Materials and Methods", "Results", "Discussion", "Conclusions"], "; ".join(top))
caps_t = sorted(int(m.group(1)) for p in doc_pars for m in [re.match(r"^Table (\d+)\. ", p)] if m)
caps_f = sorted(int(m.group(1)) for p in doc_pars for m in [re.match(r"^Figure (\d+)\. ", p)] if m)
non_caption = "\n".join(p for p in doc_pars if not re.match(r"^(Table|Figure) \d+\. ", p))
ref_t = sorted({int(x) for x in re.findall(r"Tables? (\d+)", non_caption)})
ref_f = sorted({int(x) for x in re.findall(r"Figures? (\d+)", non_caption)})
check("tables numbered consecutively", caps_t == list(range(1, len(caps_t) + 1)), str(caps_t))
check("figures numbered consecutively", caps_f == list(range(1, len(caps_f) + 1)), str(caps_f))
check("every table is referenced and every table reference exists", ref_t == caps_t, f"captions {caps_t}, references {ref_t}")
check("every figure is referenced and every figure reference exists", ref_f == caps_f, f"captions {caps_f}, references {ref_f}")
first_t = [non_caption.find(f"Table {n}") for n in caps_t]
first_f = [min(x for x in (non_caption.find(f"Figure {n}"),) if x >= 0) for n in caps_f]
check("tables and figures are first cited in numerical order", first_t == sorted(first_t) and first_f == sorted(first_f))
check("table objects present", len(doc.tables) == len(caps_t), f"{len(doc.tables)} tables")
check("picture objects present", len(doc.inline_shapes) == len(caps_f), f"{len(doc.inline_shapes)} pictures")
sec_nums = {re.match(r"^(\d+(?:\.\d+)*)\.\s", h).group(1) for h in heads}
sec_refs = set(re.findall(r"Sections? (\d+(?:\.\d+)*)", full))
check("every Section reference points to an existing heading", sec_refs <= sec_nums, f"references {sorted(sec_refs)}")
eq_nums = {int(m.group(1)) for s, t in pars if s == "MDPI_3.9_equation" for m in [re.search(r"\((\d+)\)\s*$", t)] if m}
eq_refs = {int(x) for x in re.findall(r"\((\d+)\)", " ".join(re.findall(r"Eqs?\.\s*\(\d+\)(?:–\(\d+\))?", full)))}
check("equations numbered consecutively", sorted(eq_nums) == list(range(1, len(eq_nums) + 1)), f"{len(eq_nums)} equations")
check("every equation reference points to a numbered equation", eq_refs <= eq_nums, f"references {sorted(eq_refs)}")
s_refs = {int(x) for x in re.findall(r"Table S(\d+)", "\n".join(p for p in doc_pars if not re.match(r"^Table S\d+: ", p)))}
s_list = {int(m.group(1)) for p in doc_pars for m in [re.match(r"^Table S(\d+): ", p)] if m}
check("every supplementary table cited is listed, and every listed one is cited", s_refs == s_list, f"{len(s_list)} listed")
refl = [int(m.group(1)) for s, t in pars if s == "MDPI_7.1_References" for m in [re.match(r"^(\d+)\. ", t)] if m]
cit = {int(x) for grp in re.findall(r"\[(\d+(?:[,–]\d+)*)\]", non_ref)
       for part in grp.split(",") for x in range(int(part.split("–")[0]), int(part.split("–")[-1]) + 1)}
check("references numbered in order of first citation, none uncited, none missing",
      refl == list(range(1, len(refl) + 1)) and cit == set(refl), f"{len(refl)} references")
order = [int(x.split("–")[0].split(",")[0]) for x in re.findall(r"\[(\d+(?:[,–]\d+)*)\]", non_ref)]
seen, mx, ok_order = set(), 0, True
for grp in re.findall(r"\[(\d+(?:[,–]\d+)*)\]", non_ref):
    for part in grp.split(","):
        for x in range(int(part.split("–")[0]), int(part.split("–")[-1]) + 1):
            if x not in seen:
                ok_order &= x == mx + 1
                mx = max(mx, x)
                seen.add(x)
check("first citations appear in increasing order", ok_order)
abstract = next(b["text"] for b in reg["blocks"] if b["role"] == "abstract")
words = len(re.findall(r"\S+", abstract.replace("Abstract: ", "", 1)))
check("abstract is a single paragraph within the 250-word limit", words <= 250, f"{words} words")
kw = next(b["text"] for b in reg["blocks"] if b["role"] == "keywords").replace("Keywords: ", "").split("; ")
check("three to ten keywords", 3 <= len(kw) <= 10, f"{len(kw)} keywords")
placeholders = re.findall(r"\[(?:AUTHOR TO [A-Z ]+|COMMIT HASH[^\]]*|VERIFY [A-Z]+)[^\]]*\]", full)

md = ["# NUMERIC_AUDIT_FRAMEWORK_V11B", "",
      f"Generated by `manuscript/audit_numbers_framework_v11b.py` from `{DOCX.name}` and `{REG.name}`. "
      "Numbers were extracted from the docx programmatically; CSV-backed values were re-read from their file and row "
      "and re-formatted at the printed number of decimals.", "",
      "## Summary", "",
      f"- Number tokens examined (citations, cross-reference labels, reference list, appendix list and front matter excluded): {n_tokens}",
      f"- Matched to a registered source in the same block: {n_matched}",
      f"- Single-digit literals (enumeration, mathematical constants, product names such as Landsat 8): {len(literals)}",
      f"- CSV-backed values re-read and compared at printed precision: {csv_checks}",
      f"- Flags (mismatch, unsourced number, missing text or failed check): {len(flags)}",
      f"- Abstract: {words} words (limit 250)", "",
      "## Flags", ""]
md += [f"- {f}" for f in flags] or ["None."]
md += ["", "## Structural and content checks", "", "| Check | Result | Detail |", "|---|---|---|"]
md += [f"| {n} | {r} | {d} |" for n, r, d in struct]
md += ["", "## Placeholders left for the author", ""] + [f"- {p}" for p in placeholders]
md += ["", "## Number-by-number table", ""] + lines
md += ["", "## Single-digit literals not treated as data", "", "| Block | Role | Token |", "|---|---|---|"]
md += [f"| {i} | {r} | {t} |" for i, r, t in literals]
with open(OUT, "w" if REPLACE else "x", encoding="utf-8", newline="\n") as fh:
    fh.write("\n".join(md) + "\n")
print(f"tokens {n_tokens}, matched {n_matched}, literals {len(literals)}, csv checks {csv_checks}, flags {len(flags)}, abstract words {words}")
for f in flags:
    print("FLAG:", f)
for n, r, d in struct:
    print(f"{r:5} {n} {d}")

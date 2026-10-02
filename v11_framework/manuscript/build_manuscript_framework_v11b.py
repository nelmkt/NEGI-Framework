"""Build manuscript/MANUSCRIPT_FRAMEWORK_ML_V11B_TEMPLATE_E.docx in the author's paper template (new file).

V11B differs from V11 in two ways, both on the author's instruction of 2026-10-02:
  1. Only v11 results are used. No value from the earlier surrogate run is reported.
  2. The paper is laid out in manuscript/PAPER_TEMPLATE_E.docx (styles, order of back matter, APA references).

Every number is registered by v(...) (saved CSV: file, row, column, decimals) or c(...) (constant with source),
so the numeric audit is programmatic. Headings, tables, figures, equations, supplementary tables and references
are numbered from one ordered list of blocks.

usage: python build_manuscript_framework_v11b.py <package_root> <python-docx lib dir> [draft_output_dir]
Reads only; writes two NEW files. Run with PYTHONDONTWRITEBYTECODE=1.
"""
import copy
import io
import json
import os
import re
import sys
import zipfile
from pathlib import Path

import pandas as pd

ROOT = Path(sys.argv[1]).resolve()
sys.dont_write_bytecode = True
sys.path.insert(0, sys.argv[2])
from docx import Document  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Inches, Pt  # noqa: E402

TEMPLATE = ROOT / "manuscript" / "PAPER_TEMPLATE_E.docx"
OUTDIR = Path(sys.argv[3]).resolve() if len(sys.argv) > 3 else ROOT / "manuscript"
OUT = OUTDIR / "MANUSCRIPT_FRAMEWORK_ML_V11B_TEMPLATE_E.docx"
REG = OUTDIR / "NUMERIC_REGISTRY_FRAMEWORK_V11B.json"
REPLACE = os.environ.get("NEGI_V11_REPLACE_OWN_OUTPUT") == "1"  # only for files this script itself created
for f in (OUT, REG):
    if f.exists() and not REPLACE:
        raise SystemExit(f"Refusing to overwrite {f}")

MINUS = "−"
PENDING = []
BLOCKS = []
_CSV = {}


def csv(rel):
    if rel not in _CSV:
        _CSV[rel] = pd.read_csv(ROOT / rel)
    return _CSV[rel]


def fmt(x, d):
    s = f"{abs(x):,.{d}f}"
    return (MINUS if x < 0 and float(s.replace(",", "")) != 0 else "") + s


def v(rel, row, col, d=3, plus=False, scale=1.0):
    """CSV value. row = file line number (header is row 1)."""
    raw = float(csv(rel).iloc[row - 2][col])
    text = fmt(raw * scale, d)
    if plus and raw > 0:
        text = "+" + text
    PENDING.append(dict(kind="csv", text=text, file=rel, row=row, col=col, decimals=d, scale=scale, raw=raw))
    return text


def vsum(rel, rows, col):
    raw = float(sum(csv(rel).iloc[r - 2][col] for r in rows))
    text = fmt(raw, 0)
    PENDING.append(dict(kind="csv_sum", text=text, file=rel, rows=list(rows), col=col, decimals=0, raw=raw))
    return text


def c(text, source):
    PENDING.append(dict(kind="const", text=text, source=source))
    return text


def _take():
    got = list(PENDING)
    PENDING.clear()
    return got


def H(level, title, key=None):
    BLOCKS.append(dict(t="h", level=level, title=title, key=key, reg=_take()))


def P(text, style=None):
    BLOCKS.append(dict(t="p", text=text, style=style, reg=_take()))


def BM(label, text):
    BLOCKS.append(dict(t="bm", label=label, text=text, reg=_take()))


def EQ(key, text, note):
    BLOCKS.append(dict(t="eq", key=key, text=text, note=note, reg=_take()))


def TABLE(key, caption, header, rows, widths=None):
    BLOCKS.append(dict(t="table", key=key, caption=caption, header=header, rows=rows, widths=widths, reg=_take()))


def FIG(key, image_bytes, caption, width=6.2):
    BLOCKS.append(dict(t="fig", key=key, image=image_bytes, caption=caption, width=width, reg=_take()))


HB = "tables/headline_block_v8.csv"
RR = "tables/ring_reduced_coefficients_v8_followup6.csv"
PS = "tables/pixel_subring_regression_v8_followup6.csv"
PD_ = "tables/pixel_subring_pair_differences_v8_followup7.csv"
ISO = "tables/isolation_balance_v8_followup2.csv"
COL = "tables/ring_pair_collinearity_v8_followup4.csv"
CV = "tables/model_cv.csv"
IMP = "tables/model_importance.csv"
SUP = "tables/test_support_counts.csv"
BYD = "tables/test_by_dose.csv"
SBD = "tables/test_supported_by_dose.csv"
PCE = "tables/logic_primary_class_estimates.csv"
F2E = "tables/logic_primary_figure2_estimates.csv"
CMP = "tables/gee_common_month_primary_v6.csv"
CMX = "tables/gee_common_month_paired_v6.csv"
ENE = "tables/energy.csv"
WAT = "tables/water.csv"
DSUM = "tables/decision_summary.csv"
CAND = "tables/decision_bare_places_model.csv"
LEV = "tables_revision_v11/leverage_share_check_v11.csv"
IDN = "tables_revision_v11/own_only_vs_joint_identity_v11.csv"
NF = "code/negi_original/NEGI_Framework.py"
FC = "code/src/fw_config.py"
META = "code/gee/panel_meta.json"
TPL_MS = "manuscript/MANUSCRIPT_TEMPLATE_ML_REFERENCED_V11.docx"


def nf(text, line):
    return c(text, f"{NF}:{line}")


def fc(text, line):
    return c(text, f"{FC}:{line}")


K = {
    "px": ("30", f"{FC}:34 (greened 30 m pixels of the cell's 9)"),
    "cell": ("90", f"{META} cell_m"),
    "nine": ("9", f"{FC}:35 dose_bins"),
    "ci": ("95", "code/src/fw_matching.py:71-72 (2.5th and 97.5th percentiles)"),
    "nboot": ("2,000", f"{FC}:40 n_boot"),
    "nrefit": ("200", f"{FC}:57 n_refits"),
    "folds": ("5", f"{FC}:56 cv_folds"),
    "blockdeg": ("0.05", f"{FC}:39 block_deg"),
    "y14": ("2014", f"{FC}:17 pre_years"), "y15": ("2015", f"{FC}:17 pre_years"),
    "y18": ("2018", f"{FC}:18 mid_years"), "y19": ("2019", f"{FC}:18 mid_years"),
    "y24": ("2024", f"{FC}:19 post_years"), "y25": ("2025", f"{FC}:19 post_years"),
    "d12": ("1–2", f"{FC}:35 dose_bins"), "d34": ("3–4", f"{FC}:35 dose_bins"),
    "d58": ("5–8", f"{FC}:35 dose_bins"), "d99": ("9/9", f"{FC}:35 dose_bins"),
    "tests": ("62", "author-observed local pytest result, 2026-10-02 (not run by the drafting agent)"),
}


def k(name):
    text, source = K[name]
    return c(text, source)


def cite(*keys):
    return "[@" + ",".join(keys) + "]"


# ----------------------------------------------------------------------------------------------
# References: taken verbatim from the two existing lists and from Crossref records (2026-10-02),
# then restyled to APA punctuation as the template requires. No bibliographic detail is changed.
# ----------------------------------------------------------------------------------------------
def _docx_paragraphs(path):
    import xml.etree.ElementTree as ET
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    root = ET.fromstring(zipfile.ZipFile(path).read("word/document.xml"))
    return ["".join(t.text or "" for t in ch.iter(W + "t")) for ch in root.find(W + "body") if ch.tag == W + "p"]


OGP = _docx_paragraphs(ROOT / "manuscript" / "Manuscript.docx")
TPP = _docx_paragraphs(ROOT / TPL_MS)


def _ref(paragraphs, starts):
    hit = [p for p in paragraphs if re.match(r"^\[\d+\]\s+" + re.escape(starts), p)]
    if len(hit) != 1:
        raise SystemExit(f"Reference starting {starts!r}: {len(hit)} matches")
    return re.sub(r"^\[\d+\]\s+", "", hit[0]).strip()


def apa(s):
    """Author, A.A., Author, B.B., 2020. Title. Journal 12 (3), 1-2. doi  ->
    Author, A. A., & Author, B. B. (2020). Title. Journal, 12(3), 1-2. doi   (punctuation only)."""
    s = s.strip()
    m = re.search(r", (\d{4}|n\.d\.)\. ", s)
    if m:                                             # "Authors, 2020. Title" -> "Authors (2020). Title"
        authors, rest = s[:m.start()], s[m.end():]
        s = f"{authors} ({m.group(1)}). {rest}"
    m = re.search(r" \((\d{4}|n\.d\.)\)\. ", s)
    authors, tail = s[:m.start()], s[m.start():]
    authors = re.sub(r"\b([A-Z])\.(?=[A-Z-])", r"\1. ", authors).replace(". -", ".-")   # A.B. -> A. B.
    toks = authors.split(", ")
    if len(toks) >= 4 and len(toks) % 2 == 0:         # Surname, Initials pairs: add ampersand before the last author
        authors = ", ".join(toks[:-2]) + ", & " + ", ".join(toks[-2:])
    authors = authors.replace(", & & ", ", & ")
    tail = re.sub(r" (\d+) \(([\d–-]+)\)([,.]) ", r", \1(\2)\3 ", tail)                # Journal 12 (3), -> Journal, 12(3),
    tail = re.sub(r"([A-Za-z])( \d+), ((?:e?\d|\d+–))", r"\1,\2, \3", tail, count=1) if not re.search(r", \d+\(", tail) else tail
    return (authors + tail).replace("..", ".")


_RAW = {
    "alonso2020": _ref(OGP, "Alonso"), "bowler2010": _ref(OGP, "Bowler"), "bushenkova2024": _ref(OGP, "Bushenkova"),
    "chen2016": _ref(OGP, "Chen, T."), "cheng2025": _ref(OGP, "Cheng"), "ghaffour2013": _ref(OGP, "Ghaffour"),
    "gorelick2017": _ref(OGP, "Gorelick"), "howarth2020": _ref(OGP, "Howarth"), "li2023": _ref(OGP, "Li, Y."),
    "miky2023": _ref(OGP, "Miky"), "mohamed2021": _ref(OGP, "Mohamed"), "munir2025": _ref(OGP, "Munir"),
    "ng2012": _ref(OGP, "Ng, E."), "norton2015": _ref(OGP, "Norton"), "tucker1979": _ref(OGP, "Tucker"),
    "weng2009": _ref(OGP, "Weng, Q., 2009"), "yang2025": _ref(OGP, "Yang"), "zha2003": _ref(OGP, "Zha, Y."),
    "zhao2024": _ref(OGP, "Zhao"), "patel2024": _ref(OGP, "Patel"), "roy2014": _ref(OGP, "Roy"),
    "elimelech2011": "Elimelech, M., Phillip, W.A., 2011. The future of seawater desalination: Energy, technology, "
                     "and the environment. Science 333 (6043), 712–717. https://doi.org/10.1126/science.1200488",
    "jones2019": "Jones, E., Qadir, M., van Vliet, M.T.H., Smakhtin, V., Kang, S., 2019. The state of desalination "
                 "and brine production: A global outlook. Science of the Total Environment 657, 1343–1356. "
                 "https://doi.org/10.1016/j.scitotenv.2018.12.076",
    "voutchkov2018": "Voutchkov, N., 2018. Energy use for membrane seawater desalination – current status and "
                     "trends. Desalination 431, 2–14. https://doi.org/10.1016/j.desal.2017.10.033",
    "roberts2017": "Roberts, D.R., Bahn, V., Ciuti, S., Boyce, M.S., Elith, J., Guillera-Arroita, G., Hauenstein, S., "
                   "Lahoz-Monfort, J.J., Schroder, B., Thuiller, W., Warton, D.I., Wintle, B.A., Hartig, F., "
                   "Dormann, C.F., 2017. Cross-validation strategies for data with temporal, spatial, hierarchical, "
                   "or phylogenetic structure. Ecography 40 (8), 913–929. https://doi.org/10.1111/ecog.02881",
    "ploton2020": "Ploton, P., Mortier, F., Réjou-Méchain, M., Barbier, N., Picard, N., Rossi, V., "
                  "Dormann, C., Cornu, G., Viennois, G., Bayol, N., Lyapustin, A., Gourlet-Fleury, S., "
                  "Pélissier, R., 2020. Spatial validation reveals poor predictive performance of large-scale "
                  "ecological mapping models. Nature Communications 11 (1), 4540. "
                  "https://doi.org/10.1038/s41467-020-18321-y",
    "meyer2021": "Meyer, H., Pebesma, E., 2021. Predicting into unknown space? Estimating the area of applicability "
                 "of spatial prediction models. Methods in Ecology and Evolution 12 (9), 1620–1633. "
                 "https://doi.org/10.1111/2041-210x.13650",
    "cameron2008": "Cameron, A.C., Gelbach, J.B., Miller, D.L., 2008. Bootstrap-based improvements for inference "
                   "with clustered errors. Review of Economics and Statistics 90 (3), 414–427. "
                   "https://doi.org/10.1162/rest.90.3.414",
    "angrist2009": "Angrist, J.D., Pischke, J.-S., 2009. Mostly Harmless Econometrics. Princeton University Press. "
                   "https://doi.org/10.1515/9781400829828",
    "saltelli2008": "Saltelli, A., Ratto, M., Andres, T., Campolongo, F., Cariboni, J., Gatelli, D., Saisana, M., "
                    "Tarantola, S., 2008. Global Sensitivity Analysis: The Primer. Wiley. "
                    "https://doi.org/10.1002/9780470725184 [VERIFY YEAR: the Crossref record gives 2007]",
    "wilson2014": "Wilson, G., Aruliah, D.A., Brown, C.T., Chue Hong, N.P., Davis, M., Guy, R.T., Haddock, S.H.D., "
                  "Huff, K.D., Mitchell, I.M., Plumbley, M.D., Waugh, B., White, E.P., Wilson, P., 2014. Best "
                  "practices for scientific computing. PLoS Biology 12 (1), e1001745. "
                  "https://doi.org/10.1371/journal.pbio.1001745",
    "cuthbert2022": "Cuthbert, M.O., Rau, G.C., Ekström, M., O'Carroll, D.M., Bates, A.J., 2022. Global "
                    "climate-driven trade-offs between the water retention and cooling benefits of urban greening. "
                    "Nature Communications 13 (1). https://doi.org/10.1038/s41467-022-28160-8 "
                    "[AUTHOR TO VERIFY: article number]",
}
REFS = {key: apa(val) for key, val in _RAW.items()}
# The four entries of the V11 template list are already in author-(year) form; keep them verbatim.
REFS.update({"eros2020": _ref(TPP, "Earth Resources"), "abadie2005": _ref(TPP, "Abadie"),
             "valavi2019": _ref(TPP, "Valavi"), "usgs_st": _ref(TPP, "U.S. Geological Survey")})

STAB = {
    "headline": HB, "classes": PCE, "fig2est": F2E, "ringred": RR, "subring": PS, "subpair": PD_, "iso": ISO,
    "collin": COL, "subcollin": "tables/pixel_subring_collinearity_v8_followup6.csv", "cv": CV, "imp": IMP,
    "support": SUP, "bydose": BYD, "supdose": SBD, "cmonth": CMP, "cpaired": CMX,
    "dsum": DSUM, "drank": "tables/decision_ranking.csv", "dscen": "tables/decision_scenarios.csv",
    "dpair": "tables/decision_pairwise.csv", "dcoast": "tables/decision_by_coast.csv", "cand": CAND,
    "water": WAT, "energy": ENE, "lev": LEV, "idn": IDN,
}


def S(key):
    if key not in STAB:
        raise KeyError(key)
    return f"<<S:{key}>>"


def X(kind, key):
    return f"<<{kind}:{key}>>"


if len(csv(CAND)) != 0:
    raise SystemExit("decision_bare_places_model.csv is no longer empty; the gate text must be revised")

TITLE = ("A Remote Sensing and Machine Learning Framework for Evaluating Urban Greening–Energy Trade-Offs in "
         "Desalination-Dependent Cities: Spatial Validation and Metric Diagnostics in a Jeddah Case Study")

# ==============================================================================================
# FRONT MATTER
# ==============================================================================================
BLOCKS.append(dict(t="front", reg=[]))
P("Urban greening can lower surface temperature, but in desalination-dependent cities the irrigation that sustains "
  "it has an energy cost. We present a reproducible remote sensing and machine learning framework for evaluating "
  "this trade-off, demonstrated on Jeddah. Landsat 8 composites processed in Google Earth Engine form a panel of "
  f"{k('cell')} m cells. An XGBoost model predicts land surface temperature (LST) and is scored by spatial "
  f"GroupKFold (R² = {v(CV, 2, 'r2')}, RMSE = {v(CV, 2, 'rmse_C')} °C on {v(CV, 2, 'n', 0)} non-greened "
  "cells). Before its predictions are used for greening scenarios or for the Normalized Environmental Gain Index "
  "(NEGI), they must pass two further gates: a feature-space support screen and a comparison with measured matched "
  f"contrasts. Matching {v(HB, 2, 'n_matched', 0)} greened cells outside the built-up area to never-vegetated "
  f"controls gives {v(HB, 2, 'point_C_per_pixel')} °C per greened pixel, with "
  f"{v(HB, 2, 'dominant_pooled_block_leverage_share', 1, scale=100)}% of leverage in one spatial block. An own-only "
  f"slope of {v(HB, 3, 'point_C_per_pixel')} decomposes exactly into a joint own coefficient near {MINUS}"
  f"{c('1.0', RR + ' rows 2, 6, 9 and ' + PS + ' rows 2, 8 (own coefficients -1.036 to -0.988)')} plus co-varying "
  f"neighbour greening. The model failed the gates: it had support for {vsum(SUP, (2, 3, 4, 5), 'n_supported')} of "
  f"{vsum(SUP, (2, 3, 4, 5), 'n_matched')} outside validation cells and under-predicted measured cooling by "
  f"{v(BYD, 9, 'difference_C', 2)}–{v(BYD, 12, 'difference_C', 2)} °C, so no scenario or NEGI values are "
  "reported. Absolute-LST skill does not establish counterfactual validity. We also show that NEGI's cost term is "
  "a shape assumption shared by all pathways and that its maximum is fixed by position. A hypothetical ledger "
  "converts assumed irrigation into desalination energy per degree of measured cooling. Water and energy quantities "
  "are assumed, not measured; interval coverage is unproven; and transfer to other cities is untested.",
  style="abstract")
P("urban greening; land surface temperature; remote sensing; XGBoost; spatial cross-validation; matched "
  "difference-in-differences; Normalized Environmental Gain Index; seawater desalination", style="keywords")

# ==============================================================================================
# 1 INTRODUCTION
# ==============================================================================================
H(1, "Introduction", key="intro")
H(2, "Urban Greening and the Water–Energy–Thermal Nexus", key="nexus")
P("Rapid urbanization has intensified urban heat, increasing thermal exposure and cooling demand. Urban vegetation "
  "is widely used as a nature-based mitigation strategy because shading, evapotranspiration and modification of the "
  f"surface energy balance can reduce surface temperatures {cite('weng2009', 'ng2012', 'norton2015', 'bowler2010')}. "
  "The size of these effects depends on climate, urban form, vegetation type and water availability. In arid "
  "environments, sustaining vegetation usually requires irrigation, which creates a trade-off between thermal "
  f"benefit and resource demand {cite('cuthbert2022')}.")
P("The trade-off is sharper where irrigation water comes from seawater desalination. Desalination relieves "
  f"freshwater scarcity, but producing and delivering the water requires energy {cite('ghaffour2013', 'elimelech2011')}. "
  "We use the term desalination-dependent city for a city whose municipal water supply relies predominantly on "
  "desalination; the cities we have in mind are hot, arid and coastal. Global installed desalination capacity is "
  f"concentrated in such regions {cite('jones2019')}, and reported energy use for seawater reverse osmosis is of the "
  f"order of a few kilowatt-hours per cubic metre {cite('voutchkov2018', 'elimelech2011')}. Additional vegetation in "
  "such a city may therefore improve surface thermal conditions while raising the energy needed for irrigation.")
P("Jeddah, Saudi Arabia, is the case study. Its hot-arid coastal climate, rapid urban development, freshwater "
  "scarcity and reliance on desalinated water couple urban heat mitigation with resource use "
  f"{cite('mohamed2021', 'miky2023', 'munir2025')}. Electricity demand for cooling in Saudi Arabia is closely tied to "
  f"temperature {cite('howarth2020')}.")

H(2, "Remote Sensing and Machine Learning for Urban Thermal Assessment", key="rsml")
P("Satellite remote sensing gives spatially consistent measurements of urban surfaces. The Normalized Difference "
  "Vegetation Index (NDVI) summarises vegetation greenness from red and near-infrared reflectance, the Normalized "
  "Difference Built-up Index (NDBI) summarises built-up surfaces from shortwave-infrared and near-infrared "
  "reflectance, and thermal bands give land surface temperature (LST), the radiative temperature of the ground "
  f"surface {cite('tucker1979', 'zha2003', 'weng2009', 'patel2024')}. LST is not air temperature and is not a measure "
  "of human thermal comfort.")
P("Machine-learning regressors can capture nonlinear relations and interactions in such data "
  f"{cite('alonso2020', 'li2023')}. Extreme Gradient Boosting (XGBoost), an ensemble of regression trees fitted by "
  f"gradient boosting {cite('chen2016')}, is widely used for urban thermal modelling "
  f"{cite('bushenkova2024', 'cheng2025', 'yang2025')}. Once fitted, such a model can act as a fast surrogate for "
  f"evaluating many hypothetical surface states {cite('zhao2024')}. A surrogate of this kind is empirical: it "
  "reproduces associations in the training data, and its predictions for changed inputs are not causal effects "
  "unless further assumptions hold.")
P("Two validation questions follow, and they are the computational core of this paper. First, pixels close together "
  "are statistically dependent, so a random train–test split overstates skill; validation must hold out whole "
  f"spatial blocks {cite('roberts2017', 'valavi2019', 'ploton2020')}. Second, a prediction for a modified surface is "
  "only as good as the training data near that point in feature space, which motivates an explicit support screen "
  f"in the spirit of the area of applicability {cite('meyer2021')}. Neither check tells us whether a predicted change "
  "is correct; for that, predictions must be compared with measured changes.")

H(2, "Aim, Scope and What Is Not Measured", key="aim")
P("The aim of this study is to develop a remote sensing and machine learning framework for evaluating urban "
  "greening–energy trade-offs in desalination-dependent cities. That statement needs three qualifications. "
  "(i) Trade-offs are evaluated through illustrative scenarios with an assumed irrigation-energy cost. No "
  "irrigation volumes, building energy use or desalination energy are measured anywhere in the study. (ii) Jeddah "
  "is the single demonstration site. (iii) Applicability to other desalination-dependent cities is a design intent "
  "that has not been tested: the pipeline takes city-specific inputs and parameters, but the estimates, the fitted "
  "model and the scenario ledger reported here must be recalibrated for another city and do not transfer "
  f"(Section {X('sec', 'portability')}).")
P("LST and irrigation energy are physically different quantities. A predicted LST reduction is not a building-energy "
  "saving, and the framework does not convert one into the other. It defines a dimensionless index whose "
  "assumptions are written out, and it tests the predictive model against an independent, measurement-based "
  "analysis before any model-based comparison is allowed.")

H(2, "Research Questions and Contributions", key="rq")
P("The study addresses four research questions.")
P("RQ1: How well can XGBoost predict urban LST from satellite-derived surface predictors under spatially structured "
  "validation?")
P("RQ2: How does measured LST change where vegetation was established, and do model-predicted changes agree with "
  "those measurements?")
P("RQ3: How can cooling be set against an assumed desalination-energy cost of irrigation in one normalized metric, "
  "and what are that metric's mathematical limits?")
P("RQ4: How robust are the estimates to spatial leverage, neighbour exposure, feature-space support and analysis "
  "choices?")
P("The contributions are computational and mathematical. (a) An end-to-end, scripted pipeline from satellite export "
  "to tables and figures, with seeds, hashes and tests. (b) Three validation gates that decide when model "
  "predictions may be used: spatial GroupKFold, a feature-space support screen and a matched-contrast benchmark. "
  "(c) NEGI as a defined normalized metric with every assumption stated, together with its documented limits. "
  "(d) Evidence that skill in predicting absolute LST does not give counterfactual validity. (e) Estimand and "
  "identifiability diagnostics: the exact relation between own-only and joint regression coefficients, "
  "dose-squared leverage and collinearity. (f) Tests, export reconciliation checks and a synthetic recovery test in "
  "which the true effect is known.")

# ==============================================================================================
# 2 MATERIALS AND METHODS
# ==============================================================================================
H(1, "Materials and Methods", key="methods")
H(2, "Overview of the Framework", key="overview")
P(f"Table {X('tab', 'stages')} lists the stages. The order matters: the model is trained and scored first, the "
  "matched-contrast benchmark is built independently of it, and model-based scenarios and the index are evaluated "
  "only if the model passes the gates of stage 4. Populations differ between stages and are stated with every "
  "result; their numbers are never pooled.")
TABLE("stages", "Stages of the framework, their inputs and outputs.",
      ["Stage", "Input", "Method", "Output"],
      [["1. Satellite processing", "Landsat 8 Collection 2 Level-2, Google Earth Engine",
        "Cloud masking, seasonal compositing, index calculation, cell classification",
        f"{k('cell')} m cell panel, {k('y14')}–{k('y25')}"],
       ["2. Machine-learning model", "Non-greened cells of the panel",
        "XGBoost; spatial GroupKFold; spatial refits", "LST predictor and skill metrics"],
       ["3. Matched-contrast benchmark", "Greened cells and never-vegetated controls",
        "Stratified matching, zero-intercept slopes, spatial block bootstrap, joint-ring regressions",
        "Measured contrasts and diagnostics"],
       ["4. Counterfactual gates", "Stages 2 and 3", "Feature-space support screen; model-minus-measured comparison",
        "Pass or fail for model-based use"],
       ["5. Pathways and NEGI", "Model predictions, if stage 4 is passed", "Normalized benefit minus normalized cost",
        "Index trajectories"],
       ["6. Water–energy ledger", "Measured contrasts, assumed irrigation and energy intensities",
        "Scenario accounting", "Hypothetical water and energy per degree"]],
      widths=[1.25, 1.55, 2.0, 1.45])

H(2, "Case Study and Data", key="data")
P("The data are Landsat 8 Collection 2 Level-2 surface reflectance and surface temperature "
  f"{cite('roy2014', 'eros2020')} processed in Google Earth Engine {cite('gorelick2017')} for Jeddah. The unit of "
  f"analysis is a {k('cell')} m cell made of {k('nine')} optical pixels of {k('px')} m. May–September composites "
  f"were exported for {k('y14')}, {k('y15')}, {k('y18')}, {k('y19')}, {k('y24')} and {k('y25')}. A pixel is called "
  f"greened if it was bare in both {k('y14')} and {k('y15')} and has NDVI of at least "
  f"{c('0.30', 'code/gee/export_panel.py:82')} in both {k('y24')} and {k('y25')}; a cell enters the greened class "
  "only if all nine pixels were bare at baseline. The export holds "
  f"{c('1,333', META + ' n_cells class 1')} greened cells, {c('2,988', META + ' n_cells class 2')} never-vegetated "
  f"ring cells within {c('150', 'code/gee/export_panel.py:123')} m of greening, "
  f"{c('16,363', META + ' n_cells class 3')} never-vegetated control cells more than "
  f"{c('300', 'code/gee/export_panel.py:122')} m from greening (a random "
  f"{c('10', META + ' control_fraction 0.1')}% sample) and {c('4,892', META + ' n_cells class 4')} cells from a "
  f"separate random {c('10', META + ' control_fraction 0.1')}% sample of other land. Greening is a spectral "
  f"transition; it is not verified managed irrigation. Figure {X('fig', 'map')} shows the classes.")
_tpl = zipfile.ZipFile(ROOT / TPL_MS)
FIG("map", _tpl.read("word/media/image2.png"),
    "Cell classes of the panel. (a) All four exported classes, with primary matched controls highlighted; empty "
    "areas may be outside the analysis domain or excluded by classification and sampling. (b) All classified greened "
    "cells by number of greened pixels. Greening is a spectral transition, not verified managed irrigation.")

H(2, "Machine-Learning Model and Spatial Validation", key="model")
P("An XGBoost model predicts endpoint-period LST of a cell from NDVI, NDBI, elevation, surface emissivity, distance "
  "to the coast, longitude and latitude. It is trained on the representative sample with greened cells excluded, so "
  f"it never sees a treated cell. Hyperparameters are fixed ({fc('600', 54)} trees, depth {fc('6', 54)}, learning "
  f"rate {fc('0.05', 54)}). Skill is measured by {k('folds')}-fold spatial GroupKFold on blocks of a "
  f"{k('blockdeg')}° grid (Eq. ({X('eq', 'gkf')})). Two training sets are used. Model A uses all eligible "
  f"non-greened cells; Model B also excludes training cells within {fc('300', 60)} m of greening, so that it cannot "
  f"learn from land next to treated cells. Uncertainty uses {k('nrefit')} refits on spatial block resamples.")

H(2, "Matched-Contrast Benchmark", key="validation")
P("Greened cells are compared with never-vegetated control cells in the same stratum. Strata are formed before "
  "treatment from the setting (outside the built-up area or built-up surroundings, by building cover within "
  f"{fc('500', 21)} m in {k('y15')}), a {fc('0.10', 26)}° region, distance-to-coast band, own building cover "
  f"and emissivity class; a stratum is used if it has at least {fc('3', 31)} controls. The matched contrast of a "
  f"cell (Eq. ({X('eq', 'tau')})) is summarised per greened pixel by a zero-intercept slope "
  f"(Eq. ({X('eq', 'slope')})), with intervals from a spatial block bootstrap on a {k('blockdeg')}° grid "
  f"(Eq. ({X('eq', 'boot')})). The primary match retains {v(HB, 2, 'n_matched', 0)} outside-built-up cells and "
  f"{v(PCE, 12, 'n_matched', 0)} built-up cells. A narrower match that also conditions on concurrent neighbourhood "
  "development is used as a sensitivity and for the model comparison.")

H(2, "Counterfactual Gates, Pathways and NEGI", key="gates")
P("For each matched greened cell the model predicts the LST difference between its observed surface and its "
  "baseline surface, changing NDVI and NDBI only. Gate 1 is spatially blocked skill for absolute LST. Gate 2 is the "
  f"support rule (Eq. ({X('eq', 'support')})) applied to both the baseline and the changed surface. Gate 3 is "
  "agreement between the predicted difference and the measured matched contrast within a tolerance set in advance. "
  "No tolerance was prespecified in this study, so gate 3 cannot be declared passed; the comparison is reported as "
  "a diagnostic.")
P("If the gates are passed, the framework evaluates two pathways from a common baseline, indexed by an intervention "
  "fraction f between 0 and 1. In the vegetation-only pathway NDVI is raised while NDBI is held. In the coordinated "
  "pathway NDVI and NDBI change together; in the validation package the change applied to a candidate bare cell is "
  "the paired NDVI and NDBI change observed at a real greened cell, not two independent marginal values. A "
  "coordinated pathway built this way is an observed co-variation, not an observed intervention. The pathways are "
  f"compared with NEGI (Eqs. ({X('eq', 'dt')})–({X('eq', 'negi')})).")

H(2, "Formal Definitions", key="formal")
P("Symbols are defined once here. T̂(x) is the model's predicted LST at predictor vector x; x₀ is the "
  "baseline; xⱼ(f) is pathway j at fraction f.")
EQ("dt", "ΔTⱼ(f) = T̂(x₀) − T̂(xⱼ(f))",
   "Predicted cooling; positive means cooler than baseline. Assumes the model may be evaluated at xⱼ(f), which "
   "is what the gates test.")
EQ("cost", "c(f) = fᵖ / maxₖ fₖᵖ",
   f"Normalized cost with exponent p = {nf('0.5', 749)} by default and p = {nf('1.0', 750)} as an alternative. The "
   "maximum is over the evaluated grid.")
EQ("scale", "S = maxⱼ maxₖ max(ΔTⱼ(fₖ), 0)",
   "Benefit scale: the largest positive predicted cooling over all pathways in the primary fit. S is fixed from the "
   "primary fit and reused unchanged in every refit.")
EQ("negi", "NEGIⱼ(f) = α · max(ΔTⱼ(f), 0) / S − β · c(f)",
   f"Weights α = {nf('1.0', 745)} and β = {nf('1.0', 746)} are defaults without empirical calibration. "
   "Predicted warming is clipped to zero benefit.")
P(f"Four properties follow directly from Eqs. ({X('eq', 'cost')})–({X('eq', 'negi')}).")
P("Property 1 (the cost is a shape assumption). In the code the raw cost is w₀ · η · fᵖ, with "
  f"w₀ an irrigation scaling coefficient ({nf('200', 747)}) and η a desalination energy intensity "
  f"({nf('3.5', 748)} kWh per cubic metre), divided by its own maximum at the reference w₀. At the reference "
  "setting w₀ and η cancel, leaving fᵖ / max fᵖ. The cost term is therefore the assumed shape "
  "fᵖ and not a measured or estimated energy quantity. No irrigated area, land cover or pathway-specific "
  "input enters it.")
EQ("diff", "NEGI₂(f) − NEGI₁(f) = α · [max(ΔT₂(f), 0) − max(ΔT₁(f), 0)] / S",
   "Property 2: all pathways are charged the same cost at the same f, so the cost cancels in any comparison "
   "between them.")
P("Property 3 (bounds). In the primary fit the normalized benefit lies in [0, 1], so −β ≤ NEGI ≤ "
  "α. The sign at a point is positive exactly when α · max(ΔT, 0) / S exceeds β · "
  "c(f); only the ratio α/β matters. In refits the benefit can exceed 1 because S is fixed.")
P("Property 4 (the maximum is fixed by position). At the point where the largest positive cooling occurs, the "
  "normalized benefit equals 1 and NEGI equals α − β · fᵖ, whatever the size of the cooling. "
  "Because NEGI is a composite index, any conclusion drawn from it is conditional on these choices and should be "
  f"reported with its sensitivity to them {cite('saltelli2008')}.")
P("The benchmark uses the following quantities. For greened cell i, ΔYᵢ is its LST change from the "
  "baseline years to the endpoint years, C(i) is the set of eligible never-vegetated control cells in its stratum, "
  "dᵢ is its number of greened pixels and xᵢₖ is the number of greened pixels outside the cell in "
  "distance ring k.")
EQ("tau", "τᵢ = ΔYᵢ − (1/|C(i)|) Σₖ∈C(i) ΔYₖ",
   "Matched contrast. Interpreting it as an effect assumes that, within a stratum, greened and control cells would "
   f"have had the same LST change without greening (conditional parallel trends {cite('abadie2005')}).")
EQ("slope", "b = Σᵢ dᵢ τᵢ / Σᵢ dᵢ²",
   "Zero-intercept least-squares slope of contrast on dose. It is a per-pixel summary of sampled cells, not the "
   "marginal effect of one more pixel.")
EQ("lev", "L(G) = Σᵢ∈G dᵢ² / Σᵢ dᵢ²",
   "Leverage share of a group G of cells (a spatial block or a dose class) in the denominator of the slope.")
EQ("joint", "τᵢ = β_own dᵢ + Σₖ βₖ xᵢₖ + εᵢ",
   "Joint zero-intercept model with own dose and external ring counts. The split between own and external terms "
   "is descriptive; it is not causally identified.")
EQ("ovb", "b = β_own + Σₖ βₖ γₖ,   γₖ = Σᵢ dᵢ xᵢₖ / Σᵢ dᵢ²",
   f"Exact least-squares identity on a fixed set of cells (the omitted-variable formula {cite('angrist2009')}). "
   "γₖ is the zero-intercept slope of ring-k count on own dose.")
EQ("gkf", "R² = 1 − Σᵢ (yᵢ − ŷᵢ)² / Σᵢ (yᵢ − ȳ)²,   "
          "ŷᵢ from the model trained without block g(i)",
   "Spatial GroupKFold: blocks are partitioned into folds, and every prediction comes from a model that saw no "
   "cell of the same block. Adjacent blocks can still share boundaries.")
EQ("support", "D(x) = (1/k) Σₘ₌₁..ₖ ‖z(x) − z(x₍ₘ₎)‖;   supported if D(x) ≤ q",
   f"Support rule: z standardizes each feature by the training mean and standard deviation; x₍ₘ₎ are the "
   f"k = {fc('5', 58)} nearest training cells (a cell's own row excluded); q is the {fc('95', 59)}th percentile of "
   "D among training cells. Support says the point resembles training data; it does not say the prediction is right.")
EQ("boot", "w ∼ Multinomial(B; 1/B, …, 1/B);   interval = [q₀.₀₂₅, q₀.₉₇₅] of the re-estimates",
   "Spatial block bootstrap: B blocks are resampled with replacement, every cell takes its block's weight, and "
   f"matching and estimation are repeated ({k('nboot')} draws). Coverage of the percentile interval is not "
   f"established here; with few or unequal clusters it can be poor {cite('cameron2008')}.")

H(2, "Hypothetical Water and Energy Ledger", key="ledger_m")
P("The ledger is scenario accounting, not measurement. Annual irrigation depth is reference evapotranspiration "
  f"({c('2,264', META + ' reference_et_mm_per_year 2024 = 2264.39')} mm over the greened land) multiplied by a "
  f"landscape coefficient ({fc('0.50', 84)}, {fc('0.60', 84)} or {fc('0.85', 84)}) and divided by an irrigation "
  f"efficiency ({fc('0.90', 85)}, {fc('0.75', 85)} or {fc('0.60', 85)}); the central case gives "
  f"{v(WAT, 3, 'depth_m', 2)} m per year. Energy is that volume multiplied by a source intensity: "
  f"{fc('2.5', 87)}–{fc('4.0', 87)} kWh per cubic metre for seawater reverse osmosis {cite('voutchkov2018')}. "
  "Water per degree is the assumed annual volume on a cell's greened area divided by its measured matched contrast. "
  "The same irrigation per hectare is assumed everywhere, and whether the greened cells are irrigated at all is "
  "unverified.")

H(2, "Implementation, Tests and Reproducibility", key="impl")
P("The pipeline is a set of small Python modules for the panel, matching, model, validation and ledger, with one "
  "configuration file, plus Earth Engine export scripts. The index is implemented in a separate module of the same "
  "repository. The code follows common practices for scientific software: version identifiers, fixed seeds, "
  f"automated checks and no manual editing of outputs {cite('wilson2014')}.")
P(f"The package has {k('tests')} automated tests, which passed in the author's local run. They include a synthetic "
  "recovery test: a simulated panel with a known per-pixel effect and deliberate confounding, on which the matching "
  "estimator must recover the truth and a deliberately wrong model must fail. Export reconciliation checks compare "
  f"independent pixel-count exports cell by cell; for the {c('1,333', META + ' n_cells class 1')} greened cells they "
  "report no discrepancy in identifiers, own counts or nested ring counts. Two algebraic checks were added for "
  f"this paper ({S('lev')} and {S('idn')}), and every number in the text is generated from its source table by the "
  "manuscript build script.")

# ==============================================================================================
# 3 RESULTS
# ==============================================================================================
H(1, "Results", key="results")
H(2, f"Matched Contrasts (Primary Match: {v(HB, 2, 'n_matched', 0)} Outside and {v(PCE, 12, 'n_matched', 0)} "
     "Built-Up Cells)", key="res_matched")
P(f"Figure {X('fig', 'cooling')} shows the matched contrasts and Table {X('tab', 'classes')} the own-only slopes "
  f"(Eq. ({X('eq', 'slope')})) by dose class for outside cells ({S('headline')}, {S('classes')}).")
TABLE("classes", "Outside-built-up primary matched own-only slopes by dose class. Intervals are spatial-block "
                 "percentile intervals; the last column deletes the block with the largest pooled leverage.",
      ["Greened pixels", "Matched cells", "Blocks", "°C per pixel [interval]", "Leverage share", "Dominant block deleted"],
      [[k(n), v(HB, r, "n_matched", 0), v(HB, r, "n_blocks", 0),
        f"{v(HB, r, 'point_C_per_pixel')} [{v(HB, r, 'bootstrap_lo_C')}, {v(HB, r, 'bootstrap_hi_C')}]",
        v(HB, r, "dose_squared_leverage_share", 3), v(HB, r, "dominant_block_removed_C")]
       for n, r in (("d12", 3), ("d34", 4), ("d58", 5), ("d99", 6))],
      widths=[0.9, 0.85, 0.65, 1.9, 0.9, 1.05])
P(f"Among the {v(HB, 3, 'n_matched', 0)} outside cells with {k('d12')} greened pixels the own-only slope is "
  f"{v(HB, 3, 'point_C_per_pixel')} [{v(HB, 3, 'bootstrap_lo_C')}, {v(HB, 3, 'bootstrap_hi_C')}] °C per pixel. "
  "This slope absorbs neighbouring greening that co-varies with own dose and is not a larger per-pixel effect "
  f"(Section {X('sec', 'res_estimand')}). The other rows apply the same estimator to different populations.")
P(f"The pooled slope over all {v(HB, 2, 'n_matched', 0)} outside cells is {v(HB, 2, 'point_C_per_pixel')} °C per "
  f"greened pixel (bootstrap median {v(HB, 2, 'bootstrap_median_C')}, interval [{v(HB, 2, 'bootstrap_lo_C')}, "
  f"{v(HB, 2, 'bootstrap_hi_C')}]; block-jackknife interval [{v(HB, 2, 'jackknife_lo_C')}, "
  f"{v(HB, 2, 'jackknife_hi_C')}]). It describes sampled contiguous greening blocks, not the marginal effect of a "
  f"new small patch. One spatial block holds {v(HB, 2, 'dominant_pooled_block_leverage_share', 1, scale=100)}% of "
  f"the pooled dose-squared leverage (Eq. ({X('eq', 'lev')}); recomputed in {S('lev')}); deleting it leaves "
  f"{v(HB, 2, 'n_matched_after_removal', 0)} cells and {v(HB, 2, 'dominant_block_removed_C')} °C per pixel. The "
  f"built-up pooled slope is {v(PCE, 12, 'per_pixel_C')} [{v(PCE, 12, 'per_pixel_lo_C')}, "
  f"{v(PCE, 12, 'per_pixel_hi_C')}] on {v(PCE, 12, 'n_matched', 0)} cells.")
P(f"A placebo on {v(F2E, 4, 'n_treated_matched', 0)} late-greening outside cells, measured before they greened, "
  f"gives {v(F2E, 4, 'estimate_C')} [{v(F2E, 4, 'lo_C')}, {v(F2E, 4, 'hi_C')}] °C per pixel ({S('fig2est')}). "
  "It covers only part of the sample and an early period, so it cannot verify parallel trends for all cells.")
FIG("cooling", _tpl.read("word/media/image3.png"),
    "Primary pre-treatment-only matched temperature contrasts. (a) Per-pixel contrast by year relative to the "
    "baseline years. (b) Outside the built-up area and (c) built-up surroundings: slope and placebo rows are °C "
    "per greened optical pixel; dose and ring rows are °C per cell. The ring row is a contrast, not identified "
    "spillover. Year markers omit unexported intervening years.")

H(2, f"Estimand Diagnostics ({v(HB, 3, 'n_matched', 0)} Outside Cells with {k('d12')} Greened Pixels)",
  key="res_estimand")
P(f"Five joint-ring models (Eq. ({X('eq', 'joint')})) were fitted on the same {v(HB, 3, 'n_matched', 0)} cells. "
  f"Their own coefficients are {v(RR, 2, 'beta_C_per_pixel')}, {v(RR, 6, 'beta_C_per_pixel')}, "
  f"{v(RR, 9, 'beta_C_per_pixel')}, {v(PS, 2, 'beta_C_per_pixel')} and {v(PS, 8, 'beta_C_per_pixel')} °C per "
  "pixel for the full, far-ring-omitted, outer-rings-merged, fine-inner and coarse-inner specifications; they span "
  f"{v(RR, 6, 'beta_C_per_pixel')} to {v(RR, 9, 'beta_C_per_pixel')} ({S('ringred')}, {S('subring')}). "
  f"Table {X('tab', 'identity')} verifies Eq. ({X('eq', 'ovb')}) numerically: in every specification the own "
  f"coefficient plus the external coefficients weighted by γ reproduces the own-only slope "
  f"{v(IDN, 2, 'b_own_only')} to within rounding error ({S('idn')}).")
_rows = []
_RS = "ring bounds in metres: code/src/analyze_pixel_subrings_v8_followup6.py:27-30 and " + RR + " term names"
for label, r in (("Full: 0–90, 90–180, 180–300 m", 2), ("Far ring omitted", 5),
                 ("Outer rings merged: 90–300 m", 7), ("Fine inner: 0–30, 30–60, 60–90 m", 9),
                 ("Coarse inner: 0–30, 30–90 m", 14)):
    _rows.append([c(label, _RS) if any(ch.isdigit() for ch in label) else label,
                  v(IDN, r, "beta_own"), v(IDN, r, "sum_beta_ext_times_gamma"),
                  v(IDN, r, "rhs_beta_own_plus_sum"), v(IDN, r, "b_own_only")])
TABLE("identity", "Own-only slope as the sum of the joint own coefficient and the external terms "
                  "(°C per own pixel).",
      ["Joint specification", "β_own", "Σ βₖ γₖ", "Sum", "Own-only b"], _rows,
      widths=[2.6, 0.9, 0.9, 0.9, 0.95])
P("The external terms have negative point estimates as a group, but their allocation by distance is specification "
  "sensitive. The ring counts are strongly collinear: pairwise uncentred cosines are "
  f"{v(COL, 5, 'uncentred_cosine')}, {v(COL, 6, 'uncentred_cosine')} and {v(COL, 7, 'uncentred_cosine')}, and the "
  f"design condition number rises from {v(RR, 2, 'design_condition', 1)} for the four-term model to "
  f"{v(PS, 2, 'design_condition', 1)} for the six-term model ({S('collin')}, {S('subcollin')}). The within-"
  f"{k('cell')} m pattern is not resolved: the difference between the innermost two ring coefficients is "
  f"{v(PD_, 2, 'beta_a_minus_beta_b_C_per_pixel', plus=True)} [{v(PD_, 2, 'lo_C')}, "
  f"{v(PD_, 2, 'hi_C', plus=True)}] °C per pixel, and all three paired inner-ring intervals include zero "
  f"({S('subpair')}). No mechanism is identified: thermal-pixel blur, edge effects, neighbourhood confounding and "
  "site selection all remain possible.")
P(f"Cells with no external greening nearby give own-only slopes of {v(ISO, 2, 'isolated_slope_C_per_own_pixel')}, "
  f"{v(ISO, 3, 'isolated_slope_C_per_own_pixel')} and {v(ISO, 4, 'isolated_slope_C_per_own_pixel')} "
  f"[{v(ISO, 4, 'isolated_lo_C')}, {v(ISO, 4, 'isolated_hi_C')}] °C per pixel at exclusion radii of "
  f"{v(ISO, 2, 'radius_m', 0)}, {v(ISO, 3, 'radius_m', 0)} and {v(ISO, 4, 'radius_m', 0)} m, on "
  f"{v(ISO, 2, 'isolated_cells', 0)}, {v(ISO, 3, 'isolated_cells', 0)} and {v(ISO, 4, 'isolated_cells', 0)} cells "
  f"({S('iso')}). These are different, selected sites, so the gap from the full-sample slope has no causal reading.")

H(2, f"Predictive Skill of the Model ({v(CV, 2, 'n', 0)} Non-Greened Cells)", key="res_skill")
P(f"Gate 1 is passed. Spatial GroupKFold gives R² = {v(CV, 2, 'r2')}, RMSE = {v(CV, 2, 'rmse_C')} °C and "
  f"MAE = {v(CV, 2, 'mae_C')} °C on {v(CV, 2, 'n', 0)} non-greened cells ({S('cv')}; "
  f"Figure {X('fig', 'model')}a). Gain shares in the fitted model are {v(IMP, 3, 'gain_share', 1, scale=100)}% for "
  f"NDBI, {v(IMP, 5, 'gain_share', 1, scale=100)}% for emissivity, {v(IMP, 6, 'gain_share', 1, scale=100)}% for "
  f"coast distance, {v(IMP, 7, 'gain_share', 1, scale=100)}% and {v(IMP, 8, 'gain_share', 1, scale=100)}% for "
  f"longitude and latitude, {v(IMP, 4, 'gain_share', 1, scale=100)}% for elevation and "
  f"{v(IMP, 2, 'gain_share', 1, scale=100)}% for NDVI ({S('imp')}). NDVI, the variable that greening changes most, "
  "carries the smallest share of the model's gain.")
FIG("model", (ROOT / "figures_png" / "fig3_model_test.png").read_bytes(),
    "Model diagnostics. (a) Observed against spatially held-out predicted LST for non-greened cells, both settings "
    "pooled. (b) Outside the built-up area and (c) built-up surroundings: measured matched contrast and "
    "model-predicted contrast by dose class in the concurrent-change sensitivity match, for Model A (all "
    "non-greened training cells) and Model B (training cells near greening excluded). Most outside predictions lack "
    "feature-space support. Absolute-LST skill does not validate predicted effects.")

H(2, f"Counterfactual Gates ({vsum(SUP, (2, 3, 4, 5), 'n_matched')} Outside and "
     f"{vsum(SUP, (6, 7, 8, 9), 'n_matched')} Built-Up Sensitivity Cells)", key="res_gates")
P("Gates 2 and 3 are not passed. These checks were run on the narrower concurrent-change match of "
  f"{vsum(SUP, (2, 3, 4, 5), 'n_matched')} outside and {vsum(SUP, (6, 7, 8, 9), 'n_matched')} built-up cells; support "
  "was not computed for the primary match. Model B has feature-space support for "
  f"{vsum(SUP, (2, 3, 4, 5), 'n_supported')} of the {vsum(SUP, (2, 3, 4, 5), 'n_matched')} outside cells and for "
  f"{v(SUP, 5, 'n_supported', 0)} of {v(SUP, 5, 'n_matched', 0)} fully greened cells ({S('support')}). Its predicted "
  "cooling falls short of the measured contrast in every outside dose class, by "
  f"{v(BYD, 9, 'difference_C', 2, plus=True)}, {v(BYD, 10, 'difference_C', 2, plus=True)}, "
  f"{v(BYD, 11, 'difference_C', 2, plus=True)} and {v(BYD, 12, 'difference_C', 2, plus=True)} °C for {k('d12')}, "
  f"{k('d34')}, {k('d58')} and {k('d99')} pixels, and each interval excludes zero ({S('bydose')}; "
  f"Figure {X('fig', 'model')}b). For fully greened cells the model predicts {v(BYD, 12, 'model_C', 2)} °C "
  f"against a measured {v(BYD, 12, 'measured_C', 2)} °C. In built-up surroundings, the "
  f"{v(SBD, 6, 'n_cells', 0)} supported cells with {k('d12')} pixels give a model-minus-measured difference of "
  f"{v(SBD, 6, 'difference_C', 2, plus=True)} [{v(SBD, 6, 'difference_lo_C', 2)}, "
  f"{v(SBD, 6, 'difference_hi_C', 2, plus=True)}] °C ({S('supdose')}).")
P("The framework therefore released no model-based candidate site: the candidate table is empty "
  f"({S('cand')}). Pathway predictions and NEGI values are not reported for Jeddah, because the model whose "
  "predictions they would rest on did not pass the gates. These comparisons are diagnostics of the model; they do "
  "not corroborate the matched estimates.")

H(2, f"Calendar Sensitivity ({v(CMP, 12, 'n_matched', 0)} Cells; {v(CMX, 2, 'n_matched_same_cells', 0)} "
     "Shared Cells)", key="res_calendar")
P(f"With June–September composites the pooled outside slope is {v(CMP, 12, 'per_pixel_C')} °C per pixel on "
  f"{v(CMP, 12, 'n_matched', 0)} matched cells, against {v(CMP, 2, 'per_pixel_C')} on {v(CMP, 2, 'n_matched', 0)} "
  f"cells with May–September composites ({S('cmonth')}). On {v(CMX, 2, 'n_matched_same_cells', 0)} shared cells "
  f"the paired change is {v(CMX, 2, 'new_minus_old_C')} [{v(CMX, 2, 'paired_lo_C')}, "
  f"{v(CMX, 2, 'paired_hi_C', plus=True)}] °C per pixel ({S('cpaired')}). The classified population changes "
  "between the two panels, so this is a sensitivity check and not a general robustness claim.")

H(2, "Hypothetical Water and Energy Ledger", key="res_ledger")
P(f"Table {X('tab', 'ledger')} sets the assumed irrigation of each outside dose class against its measured matched "
  f"contrast ({S('dsum')}; alternative water and energy cases in {S('dscen')}, {S('water')} and {S('energy')}). "
  "Under the central assumptions the accounting gives between "
  f"{v(DSUM, 2, 'm3_per_degree', 0)} and {v(DSUM, 4, 'm3_per_degree', 0)} cubic metres of water per year, or "
  f"{v(DSUM, 2, 'MWh_per_degree [desalinated seawater (SWRO)]', 1)} to "
  f"{v(DSUM, 4, 'MWh_per_degree [desalinated seawater (SWRO)]', 1)} MWh per year of desalination energy, for each "
  "degree of measured cell cooling. Every figure in the table is hypothetical. The differences between rows follow "
  "from the assumption of equal irrigation per hectare and from the per-pixel slopes of "
  f"Table {X('tab', 'classes')}, which Section {X('sec', 'res_estimand')} shows are not per-pixel effects; the table "
  "is not a ranking of options.")
TABLE("ledger", "Hypothetical water and desalination energy per degree of measured matched cooling, outside the "
                "built-up area, concurrent-change match, central irrigation case. Energy uses the central seawater "
                "reverse osmosis intensity. Not a ranking.",
      ["Greened pixels", "Cells", "Measured contrast (°C per cell)", "Assumed water (m³ per year)",
       "m³ per °C", "MWh per °C"],
      [[k(n), v(DSUM, r, "cells", 0), v(DSUM, r, "cooling_C", 2), v(DSUM, r, "water_m3", 0),
        v(DSUM, r, "m3_per_degree", 0), v(DSUM, r, "MWh_per_degree [desalinated seawater (SWRO)]", 1)]
       for n, r in (("d12", 2), ("d34", 3), ("d58", 4), ("d99", 5))],
      widths=[0.9, 0.7, 1.3, 1.3, 1.0, 1.0])

# ==============================================================================================
# 4 DISCUSSION
# ==============================================================================================
H(1, "Discussion", key="discussion")
H(2, "What the Gates Show About Machine-Learning Surrogates", key="disc_surrogate")
P("RQ1 and RQ2. The model predicts absolute LST well under spatially blocked validation. That figure does not "
  "answer the question that matters for scenario analysis, which is whether a predicted change in LST after a "
  "change in NDVI and NDBI is correct. The matched benchmark answers it, and the answer is no: the model had "
  "support for under two percent of outside validation cells and under-predicted measured cooling in every outside "
  "dose class, by more as dose rose. Absolute-LST skill does not give counterfactual validity.")
P("The reason is visible in the model itself. Most of its gain comes from built-up intensity, emissivity and "
  "location, which explain where the city is hot; little comes from NDVI. A model can rank places by temperature "
  "from such features and still be wrong about what happens when one of them changes at a fixed place. The "
  "practical rule of the framework is that a surrogate's scenario predictions are used only after three gates: "
  "blocked validation, a support screen for the scenario points, and agreement with measured contrasts within a "
  "tolerance set in advance.")

H(2, "NEGI as a Metric", key="disc_negi")
P("RQ3. NEGI puts predicted cooling and an assumed cost on one dimensionless scale, and its algebra makes its "
  "limits explicit. The cost is an assumed curve common to all pathways; pathway comparisons reduce to clipped "
  "predicted cooling; predicted warming is discarded; the benefit scale depends on every pathway evaluated; the "
  "weights are uncalibrated; and the largest value is fixed by where the largest cooling falls. NEGI comparisons "
  "are therefore indicative under stated assumptions. They are not a ranking of options and not an investment "
  "result, and they are only as good as the predicted cooling that enters them.")
P("For that reason no NEGI value is reported for Jeddah. Scenario and index values given in earlier versions of "
  "this work rested on a surrogate that had not been tested against measured change; they are withdrawn. The "
  "trade-off is instead shown through the ledger, which uses measured contrasts and labels every water and energy "
  "figure as an assumption.")

H(2, "Portability and Scope", key="portability")
P("The framework is designed for desalination-dependent cities and has been run on one. Applying it elsewhere "
  "requires, as inputs, a Landsat LST and NDVI series, a land and coastline boundary, a built-up mask and a "
  "land-cover layer for the new city. The following must be recalibrated rather than copied: the analysis extent "
  "and seasonal window, NDVI thresholds for bare and vegetated pixels, donor-distance and ring thresholds, matching "
  "strata and calipers, the block size for validation and bootstrap, the definition of any sub-region such as the "
  "southern belt used here, the model features and support screen, and all scenario and cost assumptions. The "
  "analysis-side settings are centralised in one configuration file; the Earth Engine export settings, the "
  "coastline file and several population checks are not. A companion audit lists each hard-coded element. The "
  "Jeddah estimates, the fitted model and the hypothetical ledger must not be transferred to another city.")

H(2, "Limitations", key="limits")
P("RQ4. Interval coverage is unproven. The spatial block bootstrap and the block jackknife are reported, but no "
  "simulation establishes their coverage here, and one block holds "
  f"{v(HB, 2, 'dominant_pooled_block_leverage_share', 1, scale=100)}% of pooled leverage. Irrigation identity is "
  "unverified: the vegetation transition is spectral, and vegetation may be sustained by groundwater, wadi flow or "
  "discharge. Confounding, thermal blur and edge effects, and neighbourhood confounding are unresolved; planting "
  "dates and untreated-trend histories are unavailable, so parallel trends cannot be verified. The Landsat surface "
  f"temperature product has retrieval dependencies, including emissivity {cite('usgs_st')}. No effect-error "
  "tolerance was prespecified. Model support was computed for the concurrent-change match and not for the primary "
  "match. LST is not air temperature, thermal comfort or population benefit. Water and energy figures are "
  "hypothetical. The framework has been run on Jeddah only.")

# ==============================================================================================
# 5 CONCLUSIONS
# ==============================================================================================
H(1, "Conclusions", key="conclusion")
P("This paper presents a remote sensing and machine learning framework for evaluating urban greening–energy "
  "trade-offs in desalination-dependent cities, with Jeddah as the case study, and subjects it to the checks a "
  "computational method should pass before its outputs are used. The framework combines scripted satellite "
  "processing, an XGBoost model scored by spatial cross-validation, a matched-contrast benchmark, explicit gates "
  "for model-based use, a normalized index with stated assumptions and a hypothetical water–energy ledger. The "
  "main findings are methodological. Spatially blocked skill for absolute LST did not carry over to predicted "
  "change, so the gates withheld model-based scenarios and index values. The index's cost term is a shape "
  "assumption shared by all pathways, and its maximum is determined by position. A large own-only per-pixel slope "
  "decomposes exactly into a smaller joint coefficient and co-varying neighbour greening, and more than half of the "
  "pooled leverage sits in one spatial block. The evidence supports the framework as a reproducible way to ask the "
  "question and to see when a model's answer should not be trusted; it does not support a preferred greening "
  "level, a ranking of options or transfer of the Jeddah numbers to another city.")

# ==============================================================================================
# BACK MATTER (template order)
# ==============================================================================================
BM("Author Contributions:", "[AUTHOR TO COMPLETE]")
BM("Funding:", "[AUTHOR TO COMPLETE]")
BM("Informed Consent Statement:", "[AUTHOR TO COMPLETE]")
BM("Data Availability Statement:",
   "Subject to the release being published, the code, the result tables listed in Appendix A and the build scripts "
   "for this manuscript will be available at https://github.com/nelmkt/NEGI-Framework [COMMIT HASH / RELEASE TAG TO "
   "ADD AFTER PUSH]. Landsat imagery is publicly available through Google Earth Engine. [AUTHOR TO COMPLETE: access "
   "terms for the exported datasets, licence and any restrictions.]")
BM("Acknowledgments:", "[AUTHOR TO COMPLETE]")
BM("Declaration of Generative AI and AI-Assisted Technologies:",
   "[AUTHOR TO COMPLETE: list every AI tool used for code, analysis and writing]")
BM("Conflicts of Interest:", "[AUTHOR TO COMPLETE]")
BLOCKS.append(dict(t="slist", reg=[]))
BLOCKS.append(dict(t="reflist", reg=[]))

# ==============================================================================================
# PASS 1: numbering
# ==============================================================================================
num = {"sec": {}, "tab": {}, "fig": {}, "eq": {}}
hc = [0, 0, 0]
nt = nfig = ne = 0
for b in BLOCKS:
    if b["t"] == "h":
        lv = b["level"]
        hc[lv - 1] += 1
        for j in range(lv, 3):
            hc[j] = 0
        b["number"] = ".".join(str(x) for x in hc[:lv])
        num["sec"][b["key"]] = b["number"]
    elif b["t"] == "table":
        nt += 1
        b["number"] = nt
        num["tab"][b["key"]] = str(nt)
    elif b["t"] == "fig":
        nfig += 1
        b["number"] = nfig
        num["fig"][b["key"]] = str(nfig)
    elif b["t"] == "eq":
        ne += 1
        b["number"] = ne
        num["eq"][b["key"]] = str(ne)

cite_order, s_order = [], []


def all_text(b):
    if b["t"] == "h":
        return [b["title"]]
    if b["t"] in ("p", "bm"):
        return [b["text"]]
    if b["t"] == "eq":
        return [b["text"], b["note"]]
    if b["t"] == "table":
        return [b["caption"]] + [x for r in b["rows"] for x in r]
    if b["t"] == "fig":
        return [b["caption"]]
    return []


for b in BLOCKS:
    for txt in all_text(b):
        for m in re.finditer(r"\[@([^\]]+)\]|<<S:([a-z0-9]+)>>", txt):
            if m.group(1):
                for key in m.group(1).split(","):
                    if key not in REFS:
                        raise SystemExit(f"Unknown reference key {key}")
                    if key not in cite_order:
                        cite_order.append(key)
            elif m.group(2) not in s_order:
                s_order.append(m.group(2))


def compress(ns):
    ns = sorted(set(ns))
    out, i = [], 0
    while i < len(ns):
        j = i
        while j + 1 < len(ns) and ns[j + 1] == ns[j] + 1:
            j += 1
        out.append(str(ns[i]) if j == i else (f"{ns[i]}–{ns[j]}" if j > i + 1 else f"{ns[i]},{ns[j]}"))
        i = j + 1
    return "[" + ",".join(out) + "]"


def resolve(txt):
    def rep(m):
        kind, key = m.group(1), m.group(2)
        if kind == "S":
            return "Table S" + str(s_order.index(key) + 1)
        if key not in num[kind]:
            raise SystemExit(f"Unresolved cross-reference {kind}:{key}")
        return num[kind][key]
    txt = re.sub(r"<<(sec|tab|fig|eq|S):([a-z0-9_]+)>>", rep, txt)
    return re.sub(r"\[@([^\]]+)\]", lambda m: compress([cite_order.index(x) + 1 for x in m.group(1).split(",")]), txt)


# ==============================================================================================
# PASS 2: render into the template
# ==============================================================================================
doc = Document(TEMPLATE)
body = doc.element.body
for child in list(body):
    if child.tag != qn("w:sectPr"):
        body.remove(child)
TEXT_PT = 13     # the template's own body runs are set to 13 pt
TABLE_PT = 10.5  # table cells are set smaller so that six columns fit the text width


def run(par, text, size=TEXT_PT, bold=None, italic=None, sup=None):
    r = par.add_run(text)
    r.font.size = Pt(size)
    r.bold, r.italic = bold, italic
    if sup:
        r.font.superscript = True
    rpr = r._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme"):   # as in the template's own runs
        fonts.set(qn(a), "minorHAnsi")
    return r


def par(style, flush=True):
    p = doc.add_paragraph(style=style)
    if flush:
        p.paragraph_format.left_indent = 0    # the template overrides the style's left indent with 0
    return p


def no_numbering(p):
    ppr = p._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), "0")
    numid = OxmlElement("w:numId"); numid.set(qn("w:val"), "0")
    numpr.append(ilvl); numpr.append(numid)
    ppr.insert(1, numpr)


def border(parent_pr, tag, edges):
    el = OxmlElement(tag)
    for edge, sz in edges:
        e = OxmlElement(f"w:{edge}")
        for kk, vv in (("val", "single"), ("sz", str(sz)), ("space", "0"), ("color", "auto")):
            e.set(qn(f"w:{kk}"), vv)
        el.append(e)
    parent_pr.append(el)


rendered = []
abstract_text = None
for i, b in enumerate(BLOCKS):
    if b["t"] == "front":
        run(par("MDPI_1.1_article_type", flush=False), "Article", size=11)
        run(par("MDPI_1.2_title", flush=False), TITLE, size=14)
        p = par("MDPI_1.3_authornames", flush=False)
        run(p, "Nelly F. Almaktoum ", size=11)
        run(p, "1", size=11, sup=True)
        p = par("p1", flush=False)
        run(p, "1", size=11, sup=True)
        run(p, "[AUTHOR TO COMPLETE: university or institution, department, country, city; e-mail]", size=11)
        for role, text in (("article type", "Article"), ("title", TITLE), ("authors", "Nelly F. Almaktoum 1"),
                           ("affiliation", "1[AUTHOR TO COMPLETE: university or institution, department, country, city; e-mail]")):
            rendered.append((i, "front matter", text))
    elif b["t"] == "h":
        title = f"{b['number']}. {resolve(b['title'])}"
        p = par("MDPI_2.1_heading1" if b["level"] == 1 else "MDPI_2.2_heading2")
        if b["level"] == 2:
            p.paragraph_format.space_before = Pt(12)
        run(p, title)
        rendered.append((i, "heading", title))
    elif b["t"] == "p":
        text = resolve(b["text"])
        if b["style"] == "abstract":
            p = par("MDPI_1.7_abstract")
            run(p, "Abstract: ", bold=True)
            run(p, text)
            abstract_text = text
            rendered.append((i, "abstract", "Abstract: " + text))
        elif b["style"] == "keywords":
            p = par("MDPI_1.8_keywords")
            run(p, "Keywords: ", bold=True)
            run(p, text)
            par("MDPI_1.9_line")
            rendered.append((i, "keywords", "Keywords: " + text))
        else:
            run(par("MDPI_3.1_text"), text)
            rendered.append((i, "paragraph", text))
    elif b["t"] == "bm":
        text = resolve(b["text"])
        p = par("MDPI_6.2_BackMatter")
        run(p, b["label"], bold=True)
        run(p, " " + text)
        rendered.append((i, "back matter", b["label"] + " " + text))
    elif b["t"] == "eq":
        text = resolve(b["text"]) + f"     ({b['number']})"
        p = par("MDPI_3.9_equation")
        run(p, text, italic=True)
        note = resolve(b["note"])
        p = par("MDPI_3.2_text_no_indent")
        run(p, note)
        rendered.append((i, "equation", text))
        rendered.append((i, "equation note", note))
    elif b["t"] == "table":
        cap = resolve(b["caption"])
        p = par("MDPI_4.1_table_caption")
        run(p, f"Table {b['number']}.", bold=True)
        run(p, " " + cap)
        t = doc.add_table(rows=1, cols=len(b["header"]))
        tblpr = t._tbl.tblPr
        for old in tblpr.findall(qn("w:tblBorders")):
            tblpr.remove(old)
        border(tblpr, "w:tblBorders", (("top", 8), ("bottom", 8)))     # three-line table, as in the template
        jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "center"); tblpr.append(jc)
        for j, h in enumerate(b["header"]):
            cell = t.rows[0].cells[j]
            cell.paragraphs[0].style = doc.styles["MDPI_4.2_table_body"]
            run(cell.paragraphs[0], resolve(h), size=TABLE_PT, bold=True)
            border(cell._tc.get_or_add_tcPr(), "w:tcBorders", (("bottom", 4),))
        for row in b["rows"]:
            cells = t.add_row().cells
            for j, x in enumerate(row):
                cells[j].paragraphs[0].style = doc.styles["MDPI_4.2_table_body"]
                run(cells[j].paragraphs[0], resolve(x), size=TABLE_PT)
        if b["widths"]:
            for row in t.rows:
                for j, w in enumerate(b["widths"]):
                    row.cells[j].width = Inches(w)
        par("MDPI_3.1_text")
        rendered.append((i, "table caption", f"Table {b['number']}. {cap}"))
        rendered.append((i, "table cells", " | ".join(resolve(x) for x in b["header"]) + " || " +
                         " || ".join(" | ".join(resolve(x) for x in row) for row in b["rows"])))
    elif b["t"] == "fig":
        p = par("MDPI_5.2_figure", flush=False)
        p.add_run().add_picture(io.BytesIO(b["image"]), width=Inches(b["width"]))
        cap = resolve(b["caption"])
        p = par("MDPI_5.1_figure_caption")
        run(p, f"Figure {b['number']}. ", bold=True)
        run(p, cap)
        rendered.append((i, "figure caption", f"Figure {b['number']}. {cap}"))
    elif b["t"] == "slist":
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.space_before = Pt(12)
        run(p, "Appendix A", bold=True)
        rendered.append((i, "appendix heading", "Appendix A"))
        line = "Supplementary tables cited in the text and the files that hold them."
        run(par("MDPI_3.1_text"), line)
        rendered.append((i, "supplementary list", line))
        for n, key in enumerate(s_order, 1):
            line = f"Table S{n}: {Path(STAB[key]).name} (folder {Path(STAB[key]).parent.as_posix()})"
            run(par("MDPI_3.2_text_no_indent"), line)
            rendered.append((i, "supplementary list", line))
    elif b["t"] == "reflist":
        run(par("MDPI_2.1_heading1"), "References")
        rendered.append((i, "references heading", "References"))
        for n, key in enumerate(cite_order, 1):
            line = f"{n}. {REFS[key]}"
            p = doc.add_paragraph(style="MDPI_7.1_References")
            no_numbering(p)
            run(p, line)
            rendered.append((i, "reference", line))

doc.save(OUT)

words = len(re.findall(r"\S+", abstract_text))
registry = dict(
    docx=OUT.name, template=TEMPLATE.name, abstract_words=words, abstract_limit=250,
    sections=num["sec"], tables=num["tab"], figures=num["fig"], equations=num["eq"],
    supplementary={f"S{n}": STAB[key] for n, key in enumerate(s_order, 1)},
    references={str(n): key for n, key in enumerate(cite_order, 1)},
    references_unused=[key for key in REFS if key not in cite_order],
    blocks=[dict(index=i, role=role, text=text, numbers=BLOCKS[i]["reg"]) for i, role, text in rendered],
)
with open(REG, "w" if REPLACE else "x", encoding="utf-8") as fh:
    json.dump(registry, fh, ensure_ascii=False, indent=1)
print("wrote", OUT, OUT.stat().st_size)
print("abstract words:", words)
print("sections:", num["sec"])
print("tables:", num["tab"], "figures:", num["fig"], "equations:", len(num["eq"]))
print("supplementary:", len(s_order), "references:", len(cite_order), "unused refs:", registry["references_unused"])
